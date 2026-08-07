"""
modules/excel_exporter.py

Exports KPI results, customer-segment performance, and
decision-support insights to a formatted Excel workbook
using openpyxl.
"""

from datetime import datetime
from pathlib import Path
import re

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from modules.decision_support import Insight
from modules.domain import KPIBundle
from modules.i18n import translate


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "reports" / "kpi_report.xlsx"


# ---------------------------------------------------------------------------
# Workbook styling
# ---------------------------------------------------------------------------

TITLE_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

SECTION_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="5B9BD5",
)

SUBHEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="DDEBF7",
)

CRITICAL_FILL = PatternFill(
    fill_type="solid",
    fgColor="F4CCCC",
)

WARNING_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

INFO_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7",
)

GOOD_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAD3",
)

THIN_BORDER = Border(
    left=Side(style="thin", color="B7B7B7"),
    right=Side(style="thin", color="B7B7B7"),
    top=Side(style="thin", color="B7B7B7"),
    bottom=Side(style="thin", color="B7B7B7"),
)


SEVERITY_FILLS = {
    "critical": CRITICAL_FILL,
    "warning": WARNING_FILL,
    "info": INFO_FILL,
    "good": GOOD_FILL,
}


def _style_title(
    worksheet: Worksheet,
    title: str,
    end_column: int,
) -> None:
    """Add and format a merged worksheet title."""

    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    title_cell = worksheet.cell(
        row=1,
        column=1,
        value=title,
    )

    title_cell.fill = TITLE_FILL
    title_cell.font = Font(
        color="FFFFFF",
        bold=True,
        size=16,
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 28


def _style_header_row(
    worksheet: Worksheet,
    row_number: int,
    start_column: int,
    end_column: int,
) -> None:
    """Apply consistent formatting to a table header row."""

    for column in range(
        start_column,
        end_column + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column,
        )

        cell.fill = HEADER_FILL
        cell.font = Font(
            color="FFFFFF",
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = THIN_BORDER


def _style_data_range(
    worksheet: Worksheet,
    start_row: int,
    end_row: int,
    start_column: int,
    end_column: int,
) -> None:
    """Add borders and alignment to a worksheet data range."""

    if end_row < start_row:
        return

    for row in worksheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def _auto_fit_columns(
    worksheet: Worksheet,
    minimum_width: int = 12,
    maximum_width: int = 45,
) -> None:
    """Adjust worksheet column widths based on cell contents."""

    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        maximum_length = 0

        for cell in column_cells:
            if isinstance(cell, MergedCell):
                continue

            if cell.value is None:
                continue

            value_length = len(str(cell.value))
            maximum_length = max(
                maximum_length,
                value_length,
            )

        adjusted_width = min(
            max(maximum_length + 2, minimum_width),
            maximum_width,
        )

        worksheet.column_dimensions[
            column_letter
        ].width = adjusted_width


def _prepare_table_sheet(
    worksheet: Worksheet,
    title: str,
    headers: list[str],
) -> None:
    """Prepare a worksheet containing one main table."""

    _style_title(
        worksheet,
        title,
        len(headers),
    )

    for column, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=3,
            column=column,
            value=header,
        )

    _style_header_row(
        worksheet,
        row_number=3,
        start_column=1,
        end_column=len(headers),
    )

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = (
        f"A3:{get_column_letter(len(headers))}3"
    )


# ---------------------------------------------------------------------------
# KPI Summary
# ---------------------------------------------------------------------------

def _create_summary_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the high-level KPI summary worksheet."""

    worksheet = workbook.active
    worksheet.title = "KPI Summary"

    _style_title(
        worksheet,
        "BI KPI Summary",
        3,
    )

    headers = [
        "Category",
        "Metric",
        "Value",
    ]

    for column, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=3,
            column=column,
            value=header,
        )

    _style_header_row(
        worksheet,
        row_number=3,
        start_column=1,
        end_column=3,
    )

    revenue = kpis["revenue"]
    sales = kpis["sales"]
    products = kpis["products"]
    regions = kpis["regions"]
    trends = kpis["trends"]
    meta = kpis["meta"]

    rows = [
        (
            "Revenue",
            "Total Revenue",
            revenue["total_revenue"],
            "currency",
        ),
        (
            "Revenue",
            "Average Order Value",
            revenue["avg_order_value"],
            "currency",
        ),
        (
            "Revenue",
            "Total Won Orders",
            revenue["total_won_orders"],
            "integer",
        ),
        (
            "Revenue",
            "Current Month",
            revenue["current_month"],
            "text",
        ),
        (
            "Revenue",
            "Current Month Revenue",
            revenue["current_month_revenue"],
            "currency",
        ),
        (
            "Revenue",
            "Previous Month",
            revenue["previous_month"],
            "text",
        ),
        (
            "Revenue",
            "Previous Month Revenue",
            revenue["previous_month_revenue"],
            "currency",
        ),
        (
            "Revenue",
            "MoM Growth",
            revenue["mom_growth_pct"],
            "percentage",
        ),
        (
            "Sales",
            "Won Orders",
            sales["won_orders"],
            "integer",
        ),
        (
            "Sales",
            "Lost Orders",
            sales["lost_orders"],
            "integer",
        ),
        (
            "Sales",
            "Churned Orders",
            sales["churned_orders"],
            "integer",
        ),
        (
            "Sales",
            "Win Rate",
            sales["win_rate_pct"],
            "percentage",
        ),
        (
            "Sales",
            "Churned Order Rate",
            sales["churned_order_rate_pct"],
            "percentage",
        ),
        (
            "Sales",
            "Top Account Manager",
            sales["top_account_manager"],
            "text",
        ),
        (
            "Sales",
            "Top Account Manager Revenue",
            sales["top_am_revenue"],
            "currency",
        ),
        (
            "Product",
            "Top Product",
            products["top_product"],
            "text",
        ),
        (
            "Region",
            "Top Region",
            regions["top_region"],
            "text",
        ),
        (
            "Trend",
            "Trend Direction",
            trends["trend_direction"],
            "text",
        ),
        (
            "Trend",
            "Trend Slope",
            trends["trend_slope"],
            "number",
        ),
        (
            "Metadata",
            "Customers in Scope",
            meta["total_customers"],
            "integer",
        ),
        (
            "Metadata",
            "Active Customers",
            meta["active_customers"],
            "integer",
        ),
        (
            "Metadata",
            "Total Orders",
            meta["total_orders"],
            "integer",
        ),
        (
            "Metadata",
            "Date Range",
            meta["date_range"],
            "text",
        ),
        (
            "Metadata",
            "Generated At",
            meta["generated_at"],
            "text",
        ),
    ]

    currency_cells = []
    percentage_cells = []

    for row_number, (
        category,
        metric,
        value,
        value_type,
    ) in enumerate(
        rows,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=category,
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=metric,
        )

        value_cell = worksheet.cell(
            row=row_number,
            column=3,
            value=value,
        )

        if value_type == "currency":
            currency_cells.append(value_cell)

        elif value_type == "percentage":
            percentage_cells.append(value_cell)

        elif value_type == "integer":
            value_cell.number_format = "#,##0"

        elif value_type == "number":
            value_cell.number_format = "#,##0.00"

    final_row = 3 + len(rows)

    _style_data_range(
        worksheet,
        start_row=4,
        end_row=final_row,
        start_column=1,
        end_column=3,
    )

    for cell in currency_cells:
        cell.number_format = '$#,##0.00'

    for cell in percentage_cells:
        cell.number_format = '0.00"%"'

    previous_category = None

    for row_number in range(
        4,
        final_row + 1,
    ):
        category_cell = worksheet.cell(
            row=row_number,
            column=1,
        )

        if category_cell.value != previous_category:
            category_cell.fill = SECTION_FILL
            category_cell.font = Font(bold=True)
            previous_category = category_cell.value

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:C{final_row}"

    _auto_fit_columns(worksheet)


# ---------------------------------------------------------------------------
# Trend worksheets
# ---------------------------------------------------------------------------

def _create_monthly_revenue_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the monthly revenue worksheet."""

    worksheet = workbook.create_sheet(
        "Monthly Revenue"
    )

    headers = [
        "Month",
        "Revenue",
    ]

    _prepare_table_sheet(
        worksheet,
        "Monthly Revenue Trend",
        headers,
    )

    records = kpis["trends"]["monthly_revenue"]

    for row_number, record in enumerate(
        records,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=str(record["month"]),
        )

        revenue_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=record["revenue"],
        )

        revenue_cell.number_format = '$#,##0.00'

    if records:
        final_row = len(records) + 3

        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=2,
        )

        worksheet.auto_filter.ref = (
            f"A3:B{final_row}"
        )

    _auto_fit_columns(worksheet)


def _create_quarterly_revenue_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the quarterly revenue worksheet."""

    worksheet = workbook.create_sheet(
        "Quarterly Revenue"
    )

    headers = [
        "Quarter",
        "Revenue",
    ]

    _prepare_table_sheet(
        worksheet,
        "Quarterly Revenue Trend",
        headers,
    )

    records = kpis["trends"]["quarterly_revenue"]

    for row_number, record in enumerate(
        records,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=str(record["quarter"]),
        )

        revenue_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=record["revenue"],
        )

        revenue_cell.number_format = '$#,##0.00'

    if records:
        final_row = len(records) + 3

        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=2,
        )

        worksheet.auto_filter.ref = (
            f"A3:B{final_row}"
        )

    _auto_fit_columns(worksheet)


# ---------------------------------------------------------------------------
# Product worksheet
# ---------------------------------------------------------------------------

def _create_product_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the product performance worksheet."""

    worksheet = workbook.create_sheet(
        "Product Performance"
    )

    headers = [
        "Product",
        "Revenue",
        "Revenue Mix",
        "Units Sold",
    ]

    _prepare_table_sheet(
        worksheet,
        "Product Performance",
        headers,
    )

    products = kpis["products"]

    product_revenue = products["product_revenue"]
    product_mix = products["product_mix_pct"]
    product_units = products["product_units_sold"]

    all_products = list(
        dict.fromkeys(
            list(product_revenue.keys())
            + list(product_mix.keys())
            + list(product_units.keys())
        )
    )

    for row_number, product in enumerate(
        all_products,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=product,
        )

        revenue_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=product_revenue.get(
                product,
                0.0,
            ),
        )
        revenue_cell.number_format = '$#,##0.00'

        mix_cell = worksheet.cell(
            row=row_number,
            column=3,
            value=product_mix.get(
                product,
                0.0,
            ),
        )
        mix_cell.number_format = '0.00"%"'

        units_cell = worksheet.cell(
            row=row_number,
            column=4,
            value=product_units.get(
                product,
                0,
            ),
        )
        units_cell.number_format = "#,##0"

    if all_products:
        final_row = len(all_products) + 3

        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=4,
        )

        worksheet.auto_filter.ref = (
            f"A3:D{final_row}"
        )

    _auto_fit_columns(worksheet)


# ---------------------------------------------------------------------------
# Regional worksheet
# ---------------------------------------------------------------------------

def _create_regional_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the regional performance worksheet."""

    worksheet = workbook.create_sheet(
        "Regional Performance"
    )

    headers = [
        "Region",
        "Revenue",
        "Won Orders",
        "Average Order Value",
    ]

    _prepare_table_sheet(
        worksheet,
        "Regional Performance",
        headers,
    )

    regions = kpis["regions"]

    regional_revenue = regions["regional_revenue"]
    regional_won_orders = regions[
        "regional_won_orders"
    ]
    regional_avg_order_value = regions[
        "regional_avg_order_value"
    ]

    all_regions = list(
        dict.fromkeys(
            list(regional_revenue.keys())
            + list(regional_won_orders.keys())
            + list(regional_avg_order_value.keys())
        )
    )

    for row_number, region in enumerate(
        all_regions,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=region,
        )

        revenue_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=regional_revenue.get(
                region,
                0.0,
            ),
        )
        revenue_cell.number_format = '$#,##0.00'

        won_orders_cell = worksheet.cell(
            row=row_number,
            column=3,
            value=regional_won_orders.get(
                region,
                0,
            ),
        )
        won_orders_cell.number_format = "#,##0"

        average_cell = worksheet.cell(
            row=row_number,
            column=4,
            value=regional_avg_order_value.get(
                region,
                0.0,
            ),
        )
        average_cell.number_format = '$#,##0.00'

    if all_regions:
        final_row = len(all_regions) + 3

        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=4,
        )

        worksheet.auto_filter.ref = (
            f"A3:D{final_row}"
        )

    _auto_fit_columns(worksheet)


# ---------------------------------------------------------------------------
# Sales worksheet
# ---------------------------------------------------------------------------

def _create_sales_sheet(
    workbook: Workbook,
    kpis: KPIBundle,
) -> None:
    """Create the sales performance worksheet."""

    worksheet = workbook.create_sheet(
        "Sales Performance"
    )

    _style_title(
        worksheet,
        "Sales Performance",
        2,
    )

    worksheet.cell(
        row=3,
        column=1,
        value="Metric",
    )
    worksheet.cell(
        row=3,
        column=2,
        value="Value",
    )

    _style_header_row(
        worksheet,
        row_number=3,
        start_column=1,
        end_column=2,
    )

    sales = kpis["sales"]

    rows = [
        (
            "Won Orders",
            sales["won_orders"],
            "integer",
        ),
        (
            "Lost Orders",
            sales["lost_orders"],
            "integer",
        ),
        (
            "Churned Orders",
            sales["churned_orders"],
            "integer",
        ),
        (
            "Win Rate",
            sales["win_rate_pct"],
            "percentage",
        ),
        (
            "Churned Order Rate",
            sales["churned_order_rate_pct"],
            "percentage",
        ),
        (
            "Top Account Manager",
            sales["top_account_manager"],
            "text",
        ),
        (
            "Top Account Manager Revenue",
            sales["top_am_revenue"],
            "currency",
        ),
    ]

    for row_number, (
        metric,
        value,
        value_type,
    ) in enumerate(
        rows,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=metric,
        )

        value_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        if value_type == "currency":
            value_cell.number_format = '$#,##0.00'

        elif value_type == "percentage":
            value_cell.number_format = '0.00"%"'

        elif value_type == "integer":
            value_cell.number_format = "#,##0"

    final_row = len(rows) + 3

    _style_data_range(
        worksheet,
        start_row=4,
        end_row=final_row,
        start_column=1,
        end_column=2,
    )

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:B{final_row}"

    _auto_fit_columns(worksheet)


# ---------------------------------------------------------------------------
# Segment worksheet
# ---------------------------------------------------------------------------

def _create_segment_sheet(
    workbook: Workbook,
    segment_analysis: dict,
) -> None:
    """Create the customer-segment performance worksheet."""

    worksheet = workbook.create_sheet(
        "Segment Performance"
    )

    headers = [
        "Segment",
        "Customers in Scope",
        "Active Customers",
        "Total Orders",
        "Won Orders",
        "Lost Orders",
        "Churned Orders",
        "Revenue",
        "Share of Selected Revenue",
        "Revenue per Active Customer",
        "Average Order Value",
        "Win Rate",
        "Churned Order Rate",
        "Top Product",
    ]

    _prepare_table_sheet(
        worksheet,
        "Customer Segment Performance",
        headers,
    )

    performance = segment_analysis.get(
        "performance",
        {},
    )

    for row_number, (
        segment,
        values,
    ) in enumerate(
        performance.items(),
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=segment,
        )

        customer_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=values["customer_count"],
        )
        customer_cell.number_format = "#,##0"

        active_customer_cell = worksheet.cell(
            row=row_number,
            column=3,
            value=values["active_customers"],
        )
        active_customer_cell.number_format = "#,##0"

        total_orders_cell = worksheet.cell(
            row=row_number,
            column=4,
            value=values["total_orders"],
        )
        total_orders_cell.number_format = "#,##0"

        won_orders_cell = worksheet.cell(
            row=row_number,
            column=5,
            value=values["won_orders"],
        )
        won_orders_cell.number_format = "#,##0"

        lost_orders_cell = worksheet.cell(
            row=row_number,
            column=6,
            value=values["lost_orders"],
        )
        lost_orders_cell.number_format = "#,##0"

        churned_orders_cell = worksheet.cell(
            row=row_number,
            column=7,
            value=values["churned_orders"],
        )
        churned_orders_cell.number_format = "#,##0"

        revenue_cell = worksheet.cell(
            row=row_number,
            column=8,
            value=values["total_revenue"],
        )
        revenue_cell.number_format = '$#,##0.00'

        share_cell = worksheet.cell(
            row=row_number,
            column=9,
            value=values["share_of_selection_pct"],
        )
        share_cell.number_format = '0.00"%"'

        revenue_per_customer_cell = worksheet.cell(
            row=row_number,
            column=10,
            value=values["revenue_per_active_customer"],
        )
        revenue_per_customer_cell.number_format = '$#,##0.00'

        average_order_value_cell = worksheet.cell(
            row=row_number,
            column=11,
            value=values["avg_order_value"],
        )
        average_order_value_cell.number_format = '$#,##0.00'

        win_rate_cell = worksheet.cell(
            row=row_number,
            column=12,
            value=values["win_rate_pct"],
        )
        win_rate_cell.number_format = '0.00"%"'

        churn_rate_cell = worksheet.cell(
            row=row_number,
            column=13,
            value=values["churned_order_rate_pct"],
        )
        churn_rate_cell.number_format = '0.00"%"'

        worksheet.cell(
            row=row_number,
            column=14,
            value=values["top_product"],
        )

    if performance:
        final_row = len(performance) + 3

        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=14,
        )

        worksheet.auto_filter.ref = (
            f"A3:N{final_row}"
        )

    worksheet.freeze_panes = "A4"

    _auto_fit_columns(
        worksheet,
        minimum_width=12,
        maximum_width=24,
    )


# ---------------------------------------------------------------------------
# Decision-support worksheet
# ---------------------------------------------------------------------------

def _create_insights_sheet(
    workbook: Workbook,
    insights_by_language: dict[
        str,
        list[Insight],
    ],
) -> None:
    """Create a bilingual decision-support insights worksheet."""

    worksheet = workbook.create_sheet(
        "Decision Insights"
    )

    headers = [
        "Language",
        "Priority",
        "Severity",
        "Category",
        "Title",
        "Detail",
        "Recommended Action",
    ]

    _prepare_table_sheet(
        worksheet,
        "Decision-Support Insights",
        headers,
    )

    row_number = 4

    for language, insights in insights_by_language.items():
        language_label = (
            "English"
            if language == "en"
            else "Türkçe"
        )

        for insight in insights:
            values = [
                language_label,
                insight.priority,
                insight.severity.capitalize(),
                insight.category,
                insight.title,
                insight.detail,
                insight.action,
            ]

            for column, value in enumerate(
                values,
                start=1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column,
                    value=value,
                )

            severity_fill = SEVERITY_FILLS.get(
                insight.severity,
                INFO_FILL,
            )

            worksheet.cell(
                row=row_number,
                column=3,
            ).fill = severity_fill

            row_number += 1

    final_row = row_number - 1

    if final_row >= 4:
        _style_data_range(
            worksheet,
            start_row=4,
            end_row=final_row,
            start_column=1,
            end_column=7,
        )

        worksheet.auto_filter.ref = (
            f"A3:G{final_row}"
        )

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 10
    worksheet.column_dimensions["C"].width = 12
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 30
    worksheet.column_dimensions["F"].width = 55
    worksheet.column_dimensions["G"].width = 55

    worksheet.freeze_panes = "A4"


def _create_forecast_sheet(
    workbook: Workbook,
    revenue_forecast: dict,
) -> None:
    """Create the revenue forecast worksheet."""

    worksheet = workbook.create_sheet(
        "Revenue Forecast"
    )

    headers = [
        "Month",
        "Actual Revenue",
        "Trend Value",
        "Forecast Revenue",
    ]

    _prepare_table_sheet(
        worksheet,
        "Three-Month Revenue Forecast",
        headers,
    )

    historical = revenue_forecast["historical"]
    forecast = revenue_forecast["forecast"]
    model = revenue_forecast["model"]

    row = 4

    for record in historical:
        worksheet.cell(
            row=row,
            column=1,
            value=record["month"],
        )

        actual = worksheet.cell(
            row=row,
            column=2,
            value=record["revenue"],
        )

        trend = worksheet.cell(
            row=row,
            column=3,
            value=record["trend_value"],
        )

        actual.number_format = '$#,##0.00'
        trend.number_format = '$#,##0.00'

        row += 1

    for record in forecast:
        worksheet.cell(
            row=row,
            column=1,
            value=record["month"],
        )

        forecast_cell = worksheet.cell(
            row=row,
            column=4,
            value=record["forecast_revenue"],
        )

        forecast_cell.number_format = '$#,##0.00'
        row += 1

    data_end_row = row - 1

    if data_end_row >= 4:
        _style_data_range(
            worksheet,
            start_row=4,
            end_row=data_end_row,
            start_column=1,
            end_column=4,
        )
        worksheet.auto_filter.ref = (
            f"A3:D{data_end_row}"
        )

    row += 2

    model_rows = [
        ("Model", model["name"]),
        ("Method", model["method"]),
        ("Slope", model["slope"]),
        (
            "R²",
            (
                model["r_squared"]
                if model.get(
                    "r_squared_interpretable",
                    True,
                )
                else "N/A"
            ),
        ),
        (
            "Forecast Reliability",
            model.get(
                "reliability",
                "unavailable",
            ).title(),
        ),
        ("Trend Direction", model["direction"]),
    ]

    model_start_row = row

    for label, value in model_rows:
        worksheet.cell(
            row=row,
            column=1,
            value=label,
        )
        worksheet.cell(
            row=row,
            column=2,
            value=value,
        )
        row += 1

    _style_data_range(
        worksheet,
        start_row=model_start_row,
        end_row=row - 1,
        start_column=1,
        end_column=2,
    )

    row += 1

    assumption_header = worksheet.cell(
        row=row,
        column=1,
        value="Forecast Assumptions",
    )
    assumption_header.font = Font(
        bold=True,
    )

    row += 1
    assumptions_start_row = row

    for assumption in revenue_forecast[
        "assumptions"
    ]:
        worksheet.cell(
            row=row,
            column=1,
            value=assumption,
        )
        row += 1

    if row > assumptions_start_row:
        _style_data_range(
            worksheet,
            start_row=assumptions_start_row,
            end_row=row - 1,
            start_column=1,
            end_column=1,
        )

    _auto_fit_columns(worksheet)
    worksheet.freeze_panes = "A4"



EXCEL_TRANSLATION_KEYS = (
    "excel.title.summary",
    "excel.title.monthly",
    "excel.title.quarterly",
    "excel.title.product",
    "excel.title.region",
    "excel.title.sales",
    "excel.title.segment",
    "excel.title.forecast",
    "excel.title.insights",
    "excel.label.category",
    "excel.label.metric",
    "excel.label.value",
    "excel.label.metadata",
    "excel.label.revenue",
    "excel.label.sales",
    "excel.label.product",
    "excel.label.region",
    "excel.label.trend",
    "excel.label.total_revenue",
    "excel.label.avg_order_value",
    "excel.label.mom_growth",
    "excel.label.win_rate",
    "excel.label.churned_order_rate",
    "excel.label.customers_scope",
    "excel.label.active_customers",
    "excel.label.total_orders",
    "excel.label.date_range",
    "excel.label.generated_at",
    "excel.label.current_month",
    "excel.label.previous_month",
    "excel.label.current_month_revenue",
    "excel.label.previous_month_revenue",
    "excel.label.total_won_orders",
    "excel.label.won_orders",
    "excel.label.lost_orders",
    "excel.label.churned_orders",
    "excel.label.top_account_manager",
    "excel.label.top_account_manager_revenue",
    "excel.label.top_product",
    "excel.label.top_region",
    "excel.label.trend_direction",
    "excel.label.trend_slope",
    "excel.label.month",
    "excel.label.quarter",
    "excel.label.revenue_mix",
    "excel.label.units_sold",
    "excel.label.segment",
    "excel.label.share_selected",
    "excel.label.revenue_per_active_customer",
    "excel.label.actual_revenue",
    "excel.label.trend_value",
    "excel.label.forecast_revenue",
    "excel.label.model",
    "excel.label.method",
    "excel.label.slope",
    "excel.label.forecast_reliability",
    "excel.label.forecast_assumptions",
    "excel.label.language",
    "excel.label.severity",
    "excel.label.priority",
    "excel.label.title",
    "excel.label.detail",
    "excel.label.recommended_action",
    "excel.assumption.history",
    "excel.assumption.linear",
    "excel.assumption.seasonality",
    "excel.assumption.directional",
    "excel.assumption.missing_months",
)

EXCEL_SHEET_KEYS = {
    "KPI Summary": "excel.sheet.summary",
    "Monthly Revenue": "excel.sheet.monthly",
    "Quarterly Revenue": "excel.sheet.quarterly",
    "Product Performance": "excel.sheet.product",
    "Regional Performance": "excel.sheet.region",
    "Sales Performance": "excel.sheet.sales",
    "Segment Performance": "excel.sheet.segment",
    "Revenue Forecast": "excel.sheet.forecast",
    "Decision Insights": "excel.sheet.insights",
}


def _bilingual_text(
    key: str,
) -> str:
    """Return one English / Turkish Excel label."""

    english = translate(
        key,
        "en",
    )
    turkish = translate(
        key,
        "tr",
    )

    if english == turkish:
        return english

    return f"{english} / {turkish}"


def _make_workbook_bilingual(
    workbook: Workbook,
) -> None:
    """
    Localize all fixed analytical labels in both supported languages.

    Business data values remain unchanged. Decision insights already
    contain separate English and Turkish records.
    """

    english_to_key = {
        translate(key, "en"): key
        for key in EXCEL_TRANSLATION_KEYS
    }

    for key in (
        "trend.direction.growing",
        "trend.direction.declining",
        "trend.direction.stable",
        "forecast.reliability.unavailable",
        "forecast.reliability.low",
        "forecast.reliability.moderate",
        "forecast.reliability.high",
    ):
        english_to_key[
            translate(key, "en")
        ] = key

    for worksheet in workbook.worksheets:
        original_title = worksheet.title

        if original_title in EXCEL_SHEET_KEYS:
            key = EXCEL_SHEET_KEYS[
                original_title
            ]
            english = translate(
                key,
                "en",
            )
            turkish = translate(
                key,
                "tr",
            )
            worksheet.title = (
                english
                if english == turkish
                else f"{english} - {turkish}"
            )[:31]

        for row in worksheet.iter_rows():
            for cell in row:
                if not isinstance(
                    cell.value,
                    str,
                ):
                    continue

                key = english_to_key.get(
                    cell.value
                )

                if key is not None:
                    cell.value = _bilingual_text(
                        key
                    )


def generate_excel_report(
    kpis: KPIBundle,
    insights_by_language: dict[
        str,
        list[Insight],
    ],
    segment_analysis: dict,
    revenue_forecast: dict,
    output_path: str | Path | None = None,
) -> str:
    """
    Generate a formatted bilingual Excel workbook containing KPI results,
    customer-segment performance, forecasts, and decision-support insights.

    Args:
        kpis:
            KPI results returned by compute_all_kpis().

        insights_by_language:
            English and Turkish Insight collections.

        segment_analysis:
            Segment-analysis results returned by
            compute_segment_analysis().

        output_path:
            Destination path of the generated Excel workbook.

    Returns:
        The generated workbook path.
    """

    if output_path is None:
        filters = kpis.get("filters", {})
        if filters:
            parts = []
            for key, value in filters.items():
                cleaned = re.sub(r"[^a-z0-9]+", "-", str(value).casefold()).strip("-")
                parts.append(f"{key}-{cleaned}")
            suffix = "_".join(parts)
        else:
            suffix = "full-dataset"
        resolved_output_path = PROJECT_ROOT / "reports" / f"kpi_report_{suffix}.xlsx"
    else:
        resolved_output_path = Path(output_path)

    resolved_output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()

    _create_summary_sheet(
        workbook,
        kpis,
    )

    _create_monthly_revenue_sheet(
        workbook,
        kpis,
    )

    _create_quarterly_revenue_sheet(
        workbook,
        kpis,
    )

    _create_product_sheet(
        workbook,
        kpis,
    )

    _create_regional_sheet(
        workbook,
        kpis,
    )

    _create_sales_sheet(
        workbook,
        kpis,
    )

    _create_segment_sheet(
        workbook,
        segment_analysis,
    )

    _create_forecast_sheet(
        workbook,
        revenue_forecast,
    )

    _create_insights_sheet(
        workbook,
        insights_by_language,
    )

    _make_workbook_bilingual(
        workbook,
    )

    workbook.properties.title = (
        "Python BI Reporting and Decision Support Report"
    )

    workbook.properties.subject = (
        "KPI, segment, and decision-support analysis"
    )

    workbook.properties.creator = (
        "Salsabeel Alfayoumi"
    )

    workbook.properties.created = datetime.now()

    workbook.save(resolved_output_path)

    print(f"  ✓ {resolved_output_path}")

    return str(resolved_output_path)